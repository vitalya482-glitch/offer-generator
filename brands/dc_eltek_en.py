from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
from typing import Any

try:
    from core.models import OfferContext
except Exception:  # pragma: no cover
    OfferContext = Any  # type: ignore

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None  # type: ignore

try:
    from num2words import num2words
except Exception:  # pragma: no cover
    num2words = None  # type: ignore

from brands import dc_eltek as base


MONTHS_EN = {
    1: "January",
    2: "February",
    3: "March",
    4: "April",
    5: "May",
    6: "June",
    7: "July",
    8: "August",
    9: "September",
    10: "October",
    11: "November",
    12: "December",
}

STATIC_EN_REPLACEMENTS = {
    "КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ": "COMMERCIAL OFFER",
    "Коммерческое предложение": "Commercial offer",
    "Кому:": "To:",
    "От:": "From:",
    "Дата:": "Date:",
    "Наименование": "Description",
    "Кол-во": "Qty",
    "Количество": "Quantity",
    "Цена за единицу": "Unit price",
    "Сумма": "Total",
    "ИТОГО": "TOTAL",
    "Итого": "Total",
    "Условия оплаты": "Payment terms",
    "Срок поставки": "Delivery time",
    "Условия поставки": "Delivery terms",
    "Монтажные работы": "Installation works",
    "Пуско-наладочные работы": "Commissioning works",
    "Пусконаладочные работы": "Commissioning works",
    "Коммерческое предложение действительно": "Commercial offer validity",
    "Гарантия": "Warranty",
    "Стоимость указана": "Price is stated",
    "С уважением": "Best regards",
    "Коммерческий директор": "Commercial Director",
    "Исполнительный директор": "Executive Director",
    "календарных дней": "calendar days",
    "тенге": "tenge",
    "евроцентов": "euro cents",
    "центов": "cents",
    "евро": "euro",
}


def _context_get(context: OfferContext | dict[str, Any], key: str, default: Any = "") -> Any:
    if isinstance(context, dict):
        return context.get(key, default)
    return getattr(context, key, default)


def _context_values(context: OfferContext | dict[str, Any]) -> dict[str, Any]:
    project_dir = str(_context_get(context, "project_dir", "") or "")
    client = str(_context_get(context, "client", "") or _context_get(context, "client_name", "") or "")
    if not client:
        client = base.extract_client_from_project_path(project_dir)
    calc_path = str(_context_get(context, "calc_path", "") or "")
    sheet_name = str(_context_get(context, "sheet_name", "") or "")
    template_path = str(_context_get(context, "template_path", "") or "") or base.find_default_dc_eltek_template()
    output_dir = str(_context_get(context, "output_dir", "") or "")
    if not output_dir and calc_path:
        output_dir = str(Path(calc_path).parent)
    if not output_dir:
        output_dir = project_dir
    return {
        "project_dir": project_dir,
        "client": client,
        "calc_path": calc_path,
        "sheet_name": sheet_name,
        "template_path": template_path,
        "output_dir": output_dir,
        "currency": str(_context_get(context, "currency", "") or "").upper().strip(),
        "signer_name": str(_context_get(context, "signer_name", "") or "Сания Санаткызы"),
        "signer_position": _translate_position(str(_context_get(context, "signer_position", "") or "Commercial Director")),
        "manager_name": str(_context_get(context, "manager_name", "") or ""),
        "manager_position": _translate_position(str(_context_get(context, "manager_position", "") or "")),
        "manager_email": str(_context_get(context, "manager_email", "") or ""),
        "manager_phone": str(_context_get(context, "manager_phone", "") or ""),
    }


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("ё", "е").lower()).strip()


def _translate_position(position: str) -> str:
    normalized = _normalize(position)
    if "коммерческий директор" in normalized:
        return "Commercial Director"
    if "исполнительный директор" in normalized:
        return "Executive Director"
    if "директор" in normalized:
        return "Director"
    return position


def _format_offer_date(dt: datetime | None = None) -> str:
    dt = dt or datetime.now()
    return f"{MONTHS_EN[dt.month]} {dt.day}, {dt.year}"


def _currency_name(currency: str) -> str:
    code = (currency or "").upper().strip()
    if code == "KZT":
        return "tenge"
    if code == "EUR":
        return "euro"
    if code == "USD":
        return "US dollars"
    return code or "currency not specified"


def _money_in_words(amount: float | int | None, currency: str) -> str:
    whole = int(round(float(amount or 0)))
    if num2words is not None:
        words = str(num2words(whole, lang="en")).replace("-", " ")
    else:
        words = str(whole)
    code = (currency or "").upper().strip()
    if code == "USD":
        return f"{words} US dollars and 00 cents"
    if code == "EUR":
        return f"{words} euro and 00 euro cents"
    if code == "KZT":
        return f"{words} tenge and 00 tiyn"
    return f"{words} {_currency_name(currency)}"


def _intro_text(summary: dict[str, Any]) -> str:
    qty = base.format_qty(summary.get("quantity_total"))
    positions_count = int(summary.get("positions_count") or 0)
    return (
        "In response to your request, we hereby submit our commercial offer for the supply "
        f"of Eltek DC power supply equipment: {positions_count} items, total quantity {qty} pcs."
    )


def _total_price_block(summary: dict[str, Any]) -> str:
    currency = str(summary.get("currency") or "")
    total = float(summary.get("total") or 0.0)
    vat_amount = float(summary.get("vat_amount") or 0.0)
    vat_label = str(summary.get("vat_label") or "0%")
    vat_text = "excluding VAT" if abs(vat_amount) < 0.01 else f"including VAT {vat_label}"
    return f"{base.format_money(total)} {_currency_name(currency)} ({_money_in_words(total, currency)}), {vat_text}."


def _currency_terms(summary: dict[str, Any]) -> str:
    currency = str(summary.get("currency") or "").upper()
    if currency == "KZT":
        return "Prices are stated in tenge."
    if currency == "EUR":
        return "Payment shall be made in tenge at the Bank CenterCredit exchange rate on the payment date."
    if currency == "USD":
        return "Payment shall be made in tenge at the bank exchange rate on the payment date."
    return "Prices are stated in the currency of this commercial offer."


def _status_from_total(total: float, included: str, not_included: str, currency: str) -> str:
    if total > 0.01:
        return f"{included}, amount {base.format_money(total)} {currency}"
    return not_included


def _special_terms_text(special_terms: dict[str, Any], currency: str) -> str:
    total = float(special_terms.get("total") or 0.0)
    rates = special_terms.get("rates") or []
    if isinstance(rates, (list, tuple, set)):
        rate_label = ", ".join(f"{float(rate):g}%" for rate in sorted({float(rate) for rate in rates}))
    else:
        rate_label = str(special_terms.get("rate_label") or "")
    if total > 0.01 or rate_label:
        if rate_label:
            return f"{rate_label}, amount {base.format_money(total)} {currency}" if total > 0.01 else rate_label
        return f"amount {base.format_money(total)} {currency}"
    return "not included"


def _build_replacements(values: dict[str, Any], parsed: dict[str, Any], offer_version: int) -> dict[str, Any]:
    summary = parsed.get("summary", {})
    meta = parsed.get("meta", {})
    currency = str(summary.get("currency") or meta.get("currency") or "")
    special_terms = meta.get("special_terms", {}) if isinstance(meta.get("special_terms", {}), dict) else {}
    installation_total = float(meta.get("installation_startup_total") or 0.0)
    inspection_total = float(meta.get("inspection_total") or 0.0)
    delivery_terms = str(meta.get("delivery_terms") or "DDP Almaty")
    installation_terms = "Installation works are included" if installation_total > 0.01 else "Installation works are not included"
    startup_terms = "Commissioning works are included" if installation_total > 0.01 else "Commissioning works are not included"
    inspection_status = _status_from_total(inspection_total, "included", "not included", currency)
    financing_terms = _special_terms_text(special_terms, currency)
    return {
        "{{offer_date}}": _format_offer_date(),
        "{{offer_version}}": str(offer_version),
        "{{client_company_full}}": values.get("client", ""),
        "{{intro_text}}": _intro_text(summary),
        "{{unit_price_header}}": f"Unit price, {_currency_name(currency)}",
        "{{total_price_header}}": f"Total, {_currency_name(currency)}",
        "{{total_label}}": "TOTAL",
        "{{grand_total}}": base.format_money(summary.get("total")),
        "{{total_price_block}}": _total_price_block(summary),
        "{{payment_terms}}": "70% advance payment, 30% after notification that the equipment is ready for shipment.",
        "{{delivery_time}}": "Delivery time shall be confirmed after order placement.",
        "{{delivery_terms}}": delivery_terms,
        "{{installation_terms}}": installation_terms,
        "{{startup_terms}}": startup_terms,
        "{{offer_validity}}": "This commercial offer is valid for 30 calendar days.",
        "{{warranty_terms}}": "Warranty: 12 months from start-up or 18 months from shipment from the factory.",
        "{{currency_terms}}": _currency_terms(summary),
        "{{signer_name}}": values.get("signer_name", ""),
        "{{signer_position}}": values.get("signer_position", ""),
        "{{manager_name}}": values.get("manager_name", ""),
        "{{manager_position}}": values.get("manager_position", ""),
        "{{manager_email}}": values.get("manager_email", ""),
        "{{manager_phone}}": values.get("manager_phone", ""),
        "{{delivery_place_terms}}": delivery_terms,
        "{{currency}}": currency,
        "{{currency_code}}": currency,
        "{{mounting_pnr_status}}": _status_from_total(installation_total, "included", "not included", currency),
        "{{installation_pnr_status}}": _status_from_total(installation_total, "included", "not included", currency),
        "{{inspection_terms}}": inspection_status,
        "{{inspection_status}}": inspection_status,
        "{{special_terms}}": financing_terms,
        "{{financing_terms}}": financing_terms,
        "{{special_financing_terms}}": financing_terms,
    }


def _paragraph_full_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs)


def _replace_literal_across_runs(paragraph, source: str, target: str) -> bool:
    if hasattr(base, "_replace_tag_across_runs"):
        return bool(base._replace_tag_across_runs(paragraph, source, target))
    full_text = _paragraph_full_text(paragraph)
    if source not in full_text or not paragraph.runs:
        return False
    paragraph.runs[0].text = full_text.replace(source, target)
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def _iter_all_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_static_texts_for_english(doc) -> None:
    for paragraph in _iter_all_paragraphs(doc):
        if not paragraph.runs:
            continue
        changed = True
        while changed:
            changed = False
            full_text = _paragraph_full_text(paragraph)
            for source, target in STATIC_EN_REPLACEMENTS.items():
                if source in full_text:
                    changed = _replace_literal_across_runs(paragraph, source, target) or changed
                    full_text = _paragraph_full_text(paragraph)


def _render_docx(template_path: str | Path, output_path: str | Path, replacements: dict[str, Any], items: list[dict[str, Any]]) -> Path:
    if Document is None:
        raise RuntimeError("Для формирования КП нужен пакет python-docx.")
    template_path = Path(template_path)
    output_path = Path(output_path)
    if not template_path.exists():
        raise FileNotFoundError(f"Шаблон КП не найден: {template_path}")
    doc = Document(str(template_path))
    base._fill_equipment_table(
        doc=doc,
        items=items,
        total_label="TOTAL",
        grand_total=str(replacements.get("{{grand_total}}", "")),
    )
    base._replace_tags(doc, replacements)
    _replace_static_texts_for_english(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path = base._get_unique_output_path(output_path)
    doc.save(str(output_path))
    return output_path


def make_offer(context: OfferContext | dict[str, Any]) -> dict[str, Any]:
    values = _context_values(context)
    calc_path = values["calc_path"]
    sheet_name = values["sheet_name"]
    template_path = values["template_path"]
    output_dir = Path(values["output_dir"] or Path(calc_path).parent)
    client = values["client"] or "DC_Eltek"
    currency_override = values.get("currency", "")

    if not calc_path:
        raise ValueError("Не выбран Excel calc.")
    if not sheet_name:
        raise ValueError("Не выбран лист для КП.")
    if not template_path:
        raise ValueError("Не выбран шаблон КП. Выберите .docx вручную или положите его в templates/dc_eltek.")

    parsed = base.read_dc_eltek_offer_items(calc_path, sheet_name, currency_override=currency_override)
    if not parsed.get("items"):
        raise ValueError("В выбранном листе Excel не найдены позиции для КП.")

    currency = str(parsed.get("summary", {}).get("currency") or "").upper().strip()
    if not currency:
        raise ValueError(
            "Валюта не указана. Укажите валюту в расчёте рядом со строками "
            "Price per unit / Total per quantity / TOTAL или выберите валюту на вкладке DC Eltek."
        )

    offer_version = base.find_next_offer_version(output_dir, client, sheet_name)
    replacements = _build_replacements(values, parsed, offer_version)
    items = base.build_offer_items(parsed)
    filename = base.build_offer_filename(client, offer_version).replace(".docx", "_EN.docx")
    output_path = _render_docx(template_path, output_dir / filename, replacements, items)
    return {
        "output_path": str(output_path),
        "language": "en",
        "items_count": len(items),
        "summary": parsed.get("summary", {}),
    }
